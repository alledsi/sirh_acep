"""Vues du module Attendance."""
from datetime import date, datetime, timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied, ValidationError
from django.db.models import Q as models_q
from django.http import Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse_lazy
from django.utils import timezone
from django.views.decorators.http import require_POST
from django.views.generic import CreateView, ListView, UpdateView, View

from apps.core.mixins import GlobalAccessRequiredMixin
from apps.employees.models import Employee

from .models import AbsenceJustification, Anomaly, TimeEntry
from .services import (
    ACTION_ARRIVAL, ACTION_BREAK_END, ACTION_BREAK_START, ACTION_DEPARTURE,
    can_punch_today, cancel_departure, get_today_entry, record_punch,
)


def _get_employee_or_none(request):
    """Récupère l'Employee lié à l'utilisateur connecté ; renvoie None sinon."""
    try:
        return request.user.employee
    except Employee.DoesNotExist:
        return None


def _no_employee_response(request):
    """Affiche un message clair quand l'utilisateur n'a pas de fiche Employee.
    Évite la boucle de redirection vers core:home."""
    return render(request, 'attendance/no_employee.html')


# ============ Dashboard agent (page d'accueil pointage) ============

@login_required
def dashboard(request):
    """Tableau de bord de l'agent — fidèle à la démo (4 stats, pointage, chart, notifs)."""
    import json

    employee = _get_employee_or_none(request)
    if not employee:
        return _no_employee_response(request)

    today_entry = get_today_entry(employee)
    can_punch, punch_message = can_punch_today(employee)

    # Stats semaine et mois
    today = timezone.localdate()
    monday = today - timedelta(days=today.weekday())
    week_entries = TimeEntry.objects.filter(
        employee=employee, work_date__gte=monday, work_date__lte=today,
    ).order_by('work_date')
    month_entries = TimeEntry.objects.filter(
        employee=employee, work_date__year=today.year, work_date__month=today.month,
    )

    def total_hours(qs):
        total = timedelta()
        days = 0
        for e in qs:
            d = e.worked_duration
            if d:
                total += d
                days += 1
        return total, days

    def fmt_hm(td: timedelta) -> str:
        """Formate une durée en 'Xh YY' (heures et minutes, sans secondes)."""
        total_minutes = int(td.total_seconds() // 60)
        h, m = divmod(total_minutes, 60)
        return f"{h}h{m:02d}"

    week_total_td, week_days = total_hours(week_entries)
    month_total_td, month_days = total_hours(month_entries)
    week_total = fmt_hm(week_total_td)
    month_total = fmt_hm(month_total_td)

    # Données du graphique hebdomadaire (5 jours Lun-Ven)
    # `minutes` est un entier = source de vérité (cohérent avec les KPIs).
    # `hours` est calculé côté front à partir des minutes pour la hauteur de barre.
    week_chart = []
    entries_by_date = {e.work_date: e for e in week_entries}
    for i in range(5):
        d = monday + timedelta(days=i)
        e = entries_by_date.get(d)
        minutes = int(e.worked_duration.total_seconds() // 60) if e and e.worked_duration else 0
        week_chart.append({
            'label': ['Lun', 'Mar', 'Mer', 'Jeu', 'Ven'][i],
            'minutes': minutes,
        })

    # Anomalies récentes (notifications)
    recent_anomalies = Anomaly.objects.filter(
        time_entry__employee=employee,
        time_entry__work_date__gte=today - timedelta(days=14),
    ).order_by('-time_entry__work_date')[:5]

    return render(request, 'attendance/dashboard.html', {
        'employee': employee,
        'today_entry': today_entry,
        'can_punch': can_punch,
        'punch_message': punch_message,
        'week_total': week_total,
        'week_days': week_days,
        'month_total': month_total,
        'month_days': month_days,
        'week_chart_data': json.dumps(week_chart),
        'recent_anomalies': recent_anomalies,
        'anomalies_today_count': today_entry.anomalies.count() if today_entry else 0,
    })


@login_required
def pointer(request):
    """Page dédiée au pointage — focus sur les 4 boutons (clock big)."""
    employee = _get_employee_or_none(request)
    if not employee:
        return _no_employee_response(request)
    today_entry = get_today_entry(employee)
    can_punch, punch_message = can_punch_today(employee)
    return render(request, 'attendance/pointer.html', {
        'employee': employee,
        'today_entry': today_entry,
        'can_punch': can_punch,
        'punch_message': punch_message,
    })


# ============ Actions de pointage (POST uniquement) ============

@login_required
@require_POST
def action_arrival(request):
    return _do_punch(request, ACTION_ARRIVAL, "Arrivée pointée.")


@login_required
@require_POST
def action_break_start(request):
    return _do_punch(request, ACTION_BREAK_START, "Début de pause pointé.")


@login_required
@require_POST
def action_break_end(request):
    return _do_punch(request, ACTION_BREAK_END, "Fin de pause pointée.")


@login_required
@require_POST
def action_departure(request):
    return _do_punch(request, ACTION_DEPARTURE, "Départ pointé.")


def _do_punch(request, action: str, success_msg: str):
    employee = _get_employee_or_none(request)
    if not employee:
        return _no_employee_response(request)

    can_punch, msg = can_punch_today(employee)
    if not can_punch:
        messages.error(request, msg)
        return redirect('attendance:dashboard')

    try:
        record_punch(action, employee, request)
        messages.success(request, success_msg)
    except ValidationError as e:
        messages.error(request, '; '.join(e.messages))
    next_url = request.POST.get('next') or 'attendance:dashboard'
    if next_url.startswith('/'):
        return redirect(next_url)
    return redirect(next_url)


@login_required
@require_POST
def action_cancel_departure(request):
    employee = _get_employee_or_none(request)
    if not employee:
        return _no_employee_response(request)
    try:
        cancel_departure(employee)
        messages.warning(request, "Départ annulé.")
    except ValidationError as e:
        messages.error(request, '; '.join(e.messages))
    return redirect('attendance:dashboard')


# ============ Historique ============

@login_required
def historique(request):
    """Historique des pointages : semaine / mois / année."""
    employee = _get_employee_or_none(request)
    if not employee:
        return _no_employee_response(request)

    periode = request.GET.get('periode', 'semaine')
    today = timezone.localdate()

    if periode == 'mois':
        start = today.replace(day=1)
        end = today
        label = today.strftime('%B %Y').capitalize()
    elif periode == 'annee':
        start = today.replace(month=1, day=1)
        end = today
        label = str(today.year)
    else:  # semaine
        periode = 'semaine'
        monday = today - timedelta(days=today.weekday())
        sunday = monday + timedelta(days=6)
        start = monday
        end = sunday
        label = f"Du {monday:%d/%m} au {sunday:%d/%m/%Y}"

    entries = (
        TimeEntry.objects
        .filter(employee=employee, work_date__gte=start, work_date__lte=end)
        .select_related('arrival_bureau', 'departure_bureau')
        .prefetch_related('anomalies')
        .order_by('work_date')
    )

    total = timedelta()
    days_count = 0
    for e in entries:
        d = e.worked_duration
        if d:
            total += d
            days_count += 1

    return render(request, 'attendance/historique.html', {
        'employee': employee,
        'periode': periode,
        'label': label,
        'entries': entries,
        'total': total,
        'days_count': days_count,
    })


@login_required
def time_entry_detail(request, pk):
    employee = _get_employee_or_none(request)
    if not employee:
        return _no_employee_response(request)
    entry = get_object_or_404(
        TimeEntry.objects.prefetch_related('anomalies'),
        pk=pk,
    )
    # Un agent voit uniquement ses propres pointages ; RH/DG voient tout
    if entry.employee_id != employee.id and not request.user.has_global_access:
        raise Http404()
    return render(request, 'attendance/time_entry_detail.html', {
        'entry': entry,
        'employee': entry.employee,
    })


# ============ Régularisation manuelle d'un pointage (RH/DG) ============

class TimeEntryRegularizeView(GlobalAccessRequiredMixin, View):
    """Permet à la RH de corriger manuellement un pointage en cas d'incident technique.

    Trace l'auteur et le motif de la régularisation.
    """
    template_name = 'attendance/regularize.html'

    def _get_entry_or_create(self, employee_id, target_date):
        """Récupère le TimeEntry de cet employé à cette date, ou le crée s'il n'existe pas."""
        emp = get_object_or_404(Employee, pk=employee_id)
        entry, _ = TimeEntry.objects.get_or_create(employee=emp, work_date=target_date)
        return entry

    def get(self, request):
        employee_id = request.GET.get('employee')
        date_str = request.GET.get('date')
        pk = request.GET.get('pk')

        entry = None
        if pk:
            entry = get_object_or_404(TimeEntry, pk=pk)
        elif employee_id and date_str:
            target_date = datetime.strptime(date_str, '%Y-%m-%d').date()
            entry = self._get_entry_or_create(int(employee_id), target_date)

        from apps.employees.models import Employee as Emp
        return render(request, self.template_name, {
            'entry': entry,
            'employees': Emp.objects.filter(is_active=True).select_related('user').order_by('user__matricule'),
        })

    def post(self, request):
        pk = request.POST.get('pk')
        if not pk:
            messages.error(request, 'Pointage introuvable.')
            return redirect('attendance:regularize')
        entry = get_object_or_404(TimeEntry, pk=pk)

        def parse_dt(field_name):
            v = request.POST.get(field_name, '').strip()
            if not v:
                return None
            try:
                return datetime.strptime(f'{entry.work_date} {v}', '%Y-%m-%d %H:%M')
            except ValueError:
                return None

        reason = request.POST.get('reason', '').strip()
        if not reason:
            messages.error(request, 'Le motif de régularisation est obligatoire.')
            return redirect(f"{reverse_lazy('attendance:regularize')}?pk={entry.pk}")

        entry.arrival_time = parse_dt('arrival_time') or entry.arrival_time
        entry.break_start = parse_dt('break_start') or entry.break_start
        entry.break_end = parse_dt('break_end') or entry.break_end
        entry.departure_time = parse_dt('departure_time') or entry.departure_time

        # Si vide explicitement, on remet à None
        for field in ('arrival_time', 'break_start', 'break_end', 'departure_time'):
            if request.POST.get(field, None) == '':
                setattr(entry, field, None)

        entry.is_regularized = True
        entry.regularization_reason = reason
        entry.regularized_by = request.user
        entry.regularized_at = timezone.now()
        entry.save()

        # Re-déclencher la détection d'anomalies
        from .services import detect_anomalies
        detect_anomalies(entry)

        messages.success(request, f'Pointage régularisé pour {entry.employee.user.matricule} ({entry.work_date:%d/%m/%Y}).')
        return redirect('attendance:regularize_list')


class RegularizationListView(GlobalAccessRequiredMixin, ListView):
    """Historique des pointages régularisés."""
    model = TimeEntry
    template_name = 'attendance/regularize_list.html'
    context_object_name = 'entries'
    paginate_by = 30

    def get_queryset(self):
        return (
            TimeEntry.objects.filter(is_regularized=True)
            .select_related('employee__user', 'regularized_by')
            .order_by('-regularized_at')
        )


# ============ Justifications d'absence (agent + RH) ============

class MyJustificationListView(LoginRequiredMixin, View):
    """Liste des justifications déposées par l'agent connecté."""
    template_name = 'attendance/justification_list_agent.html'

    def get(self, request):
        emp = _get_employee_or_none(request)
        if not emp:
            return _no_employee_response(request)
        justifications = AbsenceJustification.objects.filter(employee=emp).order_by('-absence_date')
        return render(request, self.template_name, {
            'justifications': justifications,
            'employee': emp,
        })


class JustificationCreateView(LoginRequiredMixin, View):
    """L'agent dépose une justification avec pièce jointe."""
    template_name = 'attendance/justification_form.html'

    def get(self, request):
        emp = _get_employee_or_none(request)
        if not emp:
            return _no_employee_response(request)
        return render(request, self.template_name, {
            'types': AbsenceJustification.TYPE_CHOICES,
            'today': timezone.localdate(),
        })

    def post(self, request):
        emp = _get_employee_or_none(request)
        if not emp:
            return _no_employee_response(request)

        absence_date = request.POST.get('absence_date')
        jtype = request.POST.get('justification_type', AbsenceJustification.TYPE_ABSENCE)
        reason = request.POST.get('reason', '').strip()
        attachment = request.FILES.get('attachment')

        if not absence_date or not reason:
            messages.error(request, 'Date et motif sont obligatoires.')
            return redirect('attendance:my_justifications_new')

        try:
            d = datetime.strptime(absence_date, '%Y-%m-%d').date()
        except ValueError:
            messages.error(request, 'Date invalide.')
            return redirect('attendance:my_justifications_new')

        AbsenceJustification.objects.create(
            employee=emp,
            absence_date=d,
            justification_type=jtype,
            reason=reason,
            attachment=attachment,
            status=AbsenceJustification.STATUS_PENDING,
        )
        messages.success(request, "Justification déposée. La RH ou votre directeur la traitera.")
        return redirect('attendance:my_justifications')


class JustificationReviewListView(GlobalAccessRequiredMixin, ListView):
    """Vue RH/DG : liste des justifications avec filtres et pagination."""
    model = AbsenceJustification
    template_name = 'attendance/justification_review_list.html'
    context_object_name = 'justifications'
    paginate_by = 50

    def _filtered_queryset(self):
        req = self.request
        qs = (
            AbsenceJustification.objects
            .select_related('employee__user', 'employee__bureau__agence', 'employee__direction', 'reviewed_by')
            .order_by('-absence_date')
        )
        statut = req.GET.get('statut', 'pending')
        if statut in ('PENDING', 'APPROVED', 'REJECTED'):
            qs = qs.filter(status=statut)
        elif statut == 'pending':
            qs = qs.filter(status=AbsenceJustification.STATUS_PENDING)
        # sinon 'all' → tous

        jtype = req.GET.get('type', '')
        if jtype in ('ABSENCE', 'LATE', 'EARLY_DEPARTURE', 'OTHER'):
            qs = qs.filter(justification_type=jtype)

        q = (req.GET.get('q') or '').strip()
        if q:
            qs = qs.filter(
                models_q(employee__user__matricule__icontains=q) |
                models_q(employee__user__first_name__icontains=q) |
                models_q(employee__user__last_name__icontains=q)
            )

        mutuelle_id = req.GET.get('mutuelle')
        if mutuelle_id:
            qs = qs.filter(employee__bureau__agence__mutuelle_id=mutuelle_id)
        agence_id = req.GET.get('agence')
        if agence_id:
            qs = qs.filter(employee__bureau__agence_id=agence_id)
        direction_id = req.GET.get('direction')
        if direction_id:
            qs = qs.filter(employee__direction_id=direction_id)

        start_str = req.GET.get('start')
        if start_str:
            try:
                start = datetime.strptime(start_str, '%Y-%m-%d').date()
                qs = qs.filter(absence_date__gte=start)
            except ValueError:
                pass
        end_str = req.GET.get('end')
        if end_str:
            try:
                end = datetime.strptime(end_str, '%Y-%m-%d').date()
                qs = qs.filter(absence_date__lte=end)
            except ValueError:
                pass
        return qs

    def get_queryset(self):
        return self._filtered_queryset()

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        req = self.request
        statut = req.GET.get('statut', 'pending')
        ctx['only_pending'] = statut == 'pending'
        ctx['filters'] = {
            'statut': statut,
            'type': req.GET.get('type', ''),
            'q': req.GET.get('q', ''),
            'mutuelle': req.GET.get('mutuelle', ''),
            'agence': req.GET.get('agence', ''),
            'direction': req.GET.get('direction', ''),
            'start': req.GET.get('start', ''),
            'end': req.GET.get('end', ''),
        }
        from apps.organization.models import Agence, Direction, Mutuelle
        ctx['mutuelles'] = Mutuelle.objects.filter(is_active=True)
        ctx['agences'] = Agence.objects.filter(is_active=True).select_related('mutuelle')
        ctx['directions'] = Direction.objects.filter(is_active=True)
        ctx['type_choices'] = AbsenceJustification.TYPE_CHOICES
        ctx['status_choices'] = AbsenceJustification.STATUS_CHOICES
        ctx['total_count'] = ctx['paginator'].count if ctx.get('paginator') else self._filtered_queryset().count()
        return ctx


class JustificationExportView(GlobalAccessRequiredMixin, View):
    """Export Excel de la liste filtrée des justifications."""

    def get(self, request):
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from django.http import HttpResponse

        # Réutilise la même logique de filtrage que la liste
        list_view = JustificationReviewListView()
        list_view.request = request
        qs = list_view._filtered_queryset()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Justifications'

        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='02564A', end_color='02564A', fill_type='solid')
        title_font = Font(bold=True, size=14, color='02564A')

        ws['A1'] = f"ACEP — Justifications d'absence ({qs.count()} enregistrement(s))"
        ws['A1'].font = title_font
        ws.merge_cells('A1:J1')
        ws['A1'].alignment = Alignment(horizontal='center')

        headers = ['Date', 'Matricule', 'Nom complet', 'Direction', 'Bureau',
                   'Type', 'Motif', 'Statut', 'Validée par', 'Note valideur']
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=col, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal='center')

        row_num = 4
        for j in qs:
            emp = j.employee
            ws.cell(row=row_num, column=1, value=j.absence_date.strftime('%d/%m/%Y'))
            ws.cell(row=row_num, column=2, value=emp.user.matricule)
            ws.cell(row=row_num, column=3, value=emp.user.get_full_name() or emp.user.matricule)
            ws.cell(row=row_num, column=4, value=emp.direction.name if emp.direction else '')
            ws.cell(row=row_num, column=5, value=emp.bureau.name if emp.bureau else '')
            ws.cell(row=row_num, column=6, value=j.get_justification_type_display())
            ws.cell(row=row_num, column=7, value=j.reason or '')
            ws.cell(row=row_num, column=8, value=j.get_status_display())
            ws.cell(row=row_num, column=9, value=(
                j.reviewed_by.get_full_name() if j.reviewed_by else ''))
            ws.cell(row=row_num, column=10, value=j.review_note or '')
            row_num += 1

        widths = [12, 12, 30, 25, 25, 16, 40, 14, 22, 30]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        from django.utils import timezone as _tz
        stamp = _tz.localtime().strftime('%Y%m%d_%H%M')
        response['Content-Disposition'] = f'attachment; filename="justifications_acep_{stamp}.xlsx"'
        wb.save(response)
        return response


class JustificationReviewView(GlobalAccessRequiredMixin, View):
    """Validation/rejet d'une justification + option de régularisation.

    Si des heures sont fournies (arrival_time, departure_time, etc.),
    on régularise le TimeEntry correspondant à la date d'absence.
    """

    def post(self, request, pk):
        justification = get_object_or_404(AbsenceJustification, pk=pk)
        action = request.POST.get('action')  # 'approve' / 'reject'
        note = request.POST.get('note', '').strip()

        if action == 'approve':
            justification.status = AbsenceJustification.STATUS_APPROVED
        elif action == 'reject':
            justification.status = AbsenceJustification.STATUS_REJECTED
        else:
            messages.error(request, "Action inconnue.")
            return redirect('attendance:justification_review_list')

        justification.reviewed_by = request.user
        justification.reviewed_at = timezone.now()
        justification.review_note = note
        justification.save()

        # --- Régularisation optionnelle (uniquement si approuvée + heures saisies) ---
        regularized = False
        if action == 'approve':
            def _parse_dt(field):
                v = (request.POST.get(field) or '').strip()
                if not v:
                    return None
                try:
                    naive = datetime.strptime(
                        f'{justification.absence_date} {v}', '%Y-%m-%d %H:%M',
                    )
                    return timezone.make_aware(naive) if timezone.is_naive(naive) else naive
                except ValueError:
                    return None

            arrival = _parse_dt('arrival_time')
            break_start = _parse_dt('break_start')
            break_end = _parse_dt('break_end')
            departure = _parse_dt('departure_time')

            if any([arrival, break_start, break_end, departure]):
                entry, _ = TimeEntry.objects.get_or_create(
                    employee=justification.employee,
                    work_date=justification.absence_date,
                )
                if arrival:
                    entry.arrival_time = arrival
                if break_start:
                    entry.break_start = break_start
                if break_end:
                    entry.break_end = break_end
                if departure:
                    entry.departure_time = departure
                entry.is_regularized = True
                entry.regularization_reason = (
                    f"Régularisation suite à justification approuvée "
                    f"({justification.get_justification_type_display()}) : "
                    f"{justification.reason[:200]}"
                )
                entry.regularized_by = request.user
                entry.regularized_at = timezone.now()
                entry.save()
                from .services import detect_anomalies
                detect_anomalies(entry)
                regularized = True

        if action == 'approve':
            if regularized:
                messages.success(request, "Justification approuvée et pointage régularisé.")
            else:
                messages.success(request, "Justification approuvée.")
        else:
            messages.warning(request, "Justification rejetée.")

        return redirect('attendance:justification_review_list')
